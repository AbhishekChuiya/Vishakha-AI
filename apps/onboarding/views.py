from django.utils import timezone
from .models import JoiningKitRequest

from datetime import datetime
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.http import HttpResponse, request
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from django.shortcuts import redirect, render
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from apps.onboarding.services.joining_kit_tickets import (
    create_joining_kit_tickets,
)
from .models import (
    DocumentType,
    Joinee,
    JoineeDocument,
    JoiningKitRequest,
)

HR_EMAIL = "niranshu.mittal@vishakha.com"
MANAGER_EMAIL = "abhishek.chuiya@vishakha.com"


def is_hr_user(user):
    return (
        user.is_authenticated
        and user.is_active
        and user.is_staff
        and user.email.strip().lower() == HR_EMAIL
    )


def is_joining_kit_manager(user):
    return (
        user.is_authenticated
        and user.is_active
        and user.email.strip().lower() == MANAGER_EMAIL
    )


def require_hr(user):
    if not is_hr_user(user):
        raise PermissionDenied(
            "Only the authorized HR user can perform this action."
        )


def require_assigned_manager(user, joining_kit_request):
    if not is_joining_kit_manager(user):
        raise PermissionDenied(
            "You are not authorized to access manager requests."
        )

    if (
        joining_kit_request.manager_email.strip().lower()
        != user.email.strip().lower()
    ):
        raise PermissionDenied(
            "This Joining Kit request is not assigned to you."
        )


@login_required
def dashboard(request):
    """HR dashboard for employee records and asset requests."""
    if is_joining_kit_manager(request.user):
        return redirect("onboarding:manager_joining_kit_dashboard")

    require_hr(request.user)

    joinees = list(Joinee.objects.all().order_by("-created_at"))

    total = len(joinees)
    requests_sent = 0
    awaiting_manager = 0
    submitted = 0

    for joinee in joinees:
        latest_request = (
            joinee.joining_kit_requests
            .order_by("-created_at")
            .first()
        )
        joinee.latest_kit_request = latest_request

        if latest_request:
            requests_sent += 1

            if latest_request.status == "pending":
                awaiting_manager += 1

            if latest_request.status in ("submitted", "in_progress"):
                submitted += 1

    return render(
        request,
        "onboarding/dashboard.html",
        {
            "joinees": joinees,
            "total": total,
            "requests_sent": requests_sent,
            "awaiting_manager": awaiting_manager,
            "submitted": submitted,
        },
    )

@login_required
@require_POST
def add_joinee(request):
    require_hr(request.user)

    full_name = request.POST.get("full_name", "").strip()
    if not full_name:
        return HttpResponse("Employee name is required.", status=400)

    employee_code = request.POST.get("employee_code", "").strip()

    if employee_code and Joinee.objects.filter(
        employee_code=employee_code
    ).exists():
        return HttpResponse(
            "This Employee Code is already assigned to another employee.",
            status=400,
        )

    joining_date = request.POST.get("joining_date") or None
    if joining_date:
        try:
            datetime.strptime(joining_date, "%Y-%m-%d")
        except ValueError:
            return HttpResponse("Invalid joining date.", status=400)

    Joinee.objects.create(
        full_name=full_name,
        employee_code=employee_code,
        personal_email=request.POST.get("personal_email", "").strip(),
        phone=request.POST.get("phone", "").strip(),
        designation=request.POST.get("designation", "").strip(),
        level=request.POST.get("level", "").strip().upper(),
        grade=request.POST.get("grade", "").strip().upper(),
        department=request.POST.get("department", "").strip(),
        plant_location=request.POST.get("plant_location", "").strip(),
        joining_date=joining_date,
        status=(
            "employee_code_generated"
            if employee_code
            else "pre_joining"
        ),
        created_by=request.user,
    )

    return redirect("onboarding:dashboard")



@login_required
@require_POST
def update_employee_code(request, joinee_id):
    require_hr(request.user)

    joinee = get_object_or_404(Joinee, pk=joinee_id)
    employee_code = request.POST.get("employee_code", "").strip()

    if not employee_code:
        return HttpResponse("Employee Code is required.", status=400)

    duplicate = Joinee.objects.filter(
        employee_code=employee_code
    ).exclude(pk=joinee.pk).exists()

    if duplicate:
        return HttpResponse(
            "This Employee Code is already assigned to another employee.",
            status=400,
        )

    joinee.employee_code = employee_code
    joinee.status = "employee_code_generated"
    joinee.save(update_fields=["employee_code", "status", "updated_at"])

    return redirect("onboarding:dashboard")

@login_required
def joinee_detail(request, joinee_id):
    joinee = get_object_or_404(Joinee, pk=joinee_id)

    prepare_joinee(joinee)

    documents = list(
        joinee.documents.select_related(
            "document_type",
            "updated_by",
        ).all()
    )

    applicable_documents = [
        doc for doc in documents
        if doc.status != "not_required"
    ]
    completed_documents = sum(
        doc.status == "verified"
        for doc in applicable_documents
    )
    total_documents = len(applicable_documents)

    progress = (
        round(completed_documents * 100 / total_documents)
        if total_documents
        else 0
    )

    return render(
        request,
        "onboarding/joinee_detail.html",
        {
            "joinee": joinee,
            "documents": documents,
            "document_status_choices": JoineeDocument.STATUS_CHOICES,
            "total_documents": total_documents,
            "completed_documents": completed_documents,
            "progress": progress,
        },
    )


@login_required
def export_excel(request):
    seed_default_documents()

    joinees = list(
        Joinee.objects.prefetch_related(
            "documents__document_type"
        ).all()
    )

    for joinee in joinees:
        prepare_joinee(joinee)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Pre-Joining Tracker"

    summary.append([
        "Onboarding Reference",
        "Full Name",
        "Personal Email",
        "Phone",
        "Designation",
        "Department",
        "Plant / Location",
        "Joining Date",
        "Employee Code",
        "Overall Status",
        "Created At",
    ])

    details = workbook.create_sheet("Document Statuses")
    details.append([
        "Onboarding Reference",
        "Full Name",
        "Employee Code",
        "Document Category",
        "Document Name",
        "Document Status",
        "Remarks",
        "Last Updated",
    ])

    for joinee in joinees:
        summary.append([
            str(joinee.onboarding_reference),
            joinee.full_name,
            joinee.personal_email,
            joinee.phone,
            joinee.designation,
            joinee.department,
            joinee.plant_location,
            joinee.joining_date.isoformat()
            if joinee.joining_date else "",
            joinee.employee_code,
            joinee.get_status_display(),
            joinee.created_at.strftime("%Y-%m-%d %H:%M"),
        ])

        for doc in joinee.documents.all():
            details.append([
                str(joinee.onboarding_reference),
                joinee.full_name,
                joinee.employee_code,
                doc.document_type.get_category_display(),
                doc.document_type.name,
                doc.get_status_display(),
                doc.remarks,
                doc.updated_at.strftime("%Y-%m-%d %H:%M"),
            ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="17365D",
            )

        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            sheet.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = min(max(max_length + 2, 14), 38)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        'attachment; filename="Vishakha_PreJoining_Tracker.xlsx"'
    )
    return response

ASSET_OPTIONS = [
    ("laptop", "Laptop"),
    ("desktop", "Desktop Computer"),
    ("monitor", "Monitor"),
    ("keyboard_mouse", "Keyboard and Mouse"),
    ("headset", "Headset"),
    ("mobile", "Mobile Phone"),
    ("sim", "SIM Card"),
    ("tablet", "Tablet"),
    ("email", "Email Account"),
    ("sap", "ERP / SAP Access"),
    ("peoplestrong", "PeopleStrong Access"),
    ("software", "Software / Licences"),
    ("access_card", "Access Card"),
    ("safety_shoes", "Safety Shoes"),
    ("helmet", "Safety Helmet"),
    ("safety_glasses", "Safety Glasses"),
    ("uniform", "Uniform"),
    ("other_ppe", "Other PPE / Safety Kit"),
    ("desk_chair", "Office Desk and Chair"),
    ("tools", "Job-Specific Tools"),
    ("other", "Other Asset"),
]
# Department-specific computer eligibility.
# "yes" means recommend a laptop by default.
# "no" means do not preselect a computer.
# "need_base" means the manager must decide.

DEPARTMENT_COMPUTER_RULES = {
    "ADMINISTRATION": "yes",
    "ENVIRONMENT HEALTH & SAFETY": "need_base",
    "ANODIZING": "no",
    "ASSEMBLY": "no",
    "PROCESS AUDIT": "yes",
    "BUSINESS DEVELOPMENT": "yes",
    "BUSINESS / OPERATIONAL EXCELLENCE": "yes",
    "CIVIL": "no",
    "CONVERTING": "no",
    "CORPORATE AFFAIRS": "yes",
    "CORPORATE COMMUNICATION": "yes",
    "HUMAN RESOURCE": "yes",
    "INFORMATION TECHNOLOGY": "yes",
    "CORPORATE PROCUREMENT": "yes",
    "TRAVEL DESK": "yes",
    "DATA ANALYTICS": "yes",
    "DEEP PROCESS": "no",
    "DESIGN": "yes",
    "DIRECTOR'S OFFICE": "yes",
    "DISPATCH": "need_base",
    "ENGINEERING": "no",
    "EXIM": "yes",
    "EXTRUSION": "no",
    "FINANCE/ACCOUNTS": "yes",
    "FURNACE/FOUNDRY": "no",
    "GRAPHICS": "yes",
    "LAMINATION": "no",
    "LEGAL": "yes",
    "LOGISTIC": "yes",
    "MACHINE SHOP": "no",
    "MANAGEMENT AUDIT/ASSURANCE SERVICES": "yes",
    "MOULDING": "no",
    "MR & DOCUMENTATION": "yes",
    "NEW PRODUCT DEVELOPMENT": "yes",
    "OPERATIONS": "yes",
    "PE": "no",
    "PLANNING/STRATEGY": "yes",
    "EXPORT": "yes",
    "POUCH": "no",
    "PRINTING": "no",
    "PROCUREMENT": "yes",
    "PRODUCTION": "no",
    "PRODUCTION PLANNING & CONTROL": "yes",
    "PROJECT PLANNING & CONTROL": "yes",
    "PROJECTS": "no",
    "SUPPLY CHAIN MANAGEMENT": "need_base",
    "PROJECTS CIVIL & STRUCTURE": "need_base",
    "QUALITY CONTROL": "no",
    "PVC": "no",
    "QAME": "no",
    "RECYCLING": "no",
    "RESEARCH DEVELOPMENT": "yes",
    "REWINDING": "no",
    "RM PROCUREMENT": "yes",
    "ROLLED GLASS": "no",
    "SALES & MARKETING": "yes",
    "SAND": "no",
    "SECRETARIAL": "yes",
    "SECURITY": "no",
    "SLITTING": "no",
    "STORES": "need_base",
    "TOOL ROOM": "no",
}


def get_recommended_assets(joinee):
    """
    Return default asset keys for a joinee.
    These are recommendations only; the manager may change them.
    """
    level = (joinee.level or "").strip().upper()
    grade = (joinee.grade or "").strip().upper()
    department = " ".join(
        (joinee.department or "").upper().replace("\u00a0", " ").split()
    )

    computer_rule = DEPARTMENT_COMPUTER_RULES.get(
        department,
        "need_base",
    )

    recommended = []

    # L2, L3 and L4-D1 have standard laptop and email defaults.
    if level in ("L2", "L3") or (
        level == "L4" and grade == "D1"
    ):
        recommended.extend(["laptop", "email"])

    # L4-D2/D3 and all L5/L6 use department computer rules.
    elif (
        (level == "L4" and grade in ("D2", "D3"))
        or level in ("L5", "L6")
    ):
        if computer_rule == "yes":
            recommended.append("laptop")

        # Email is only auto-recommended for Support departments
        # in L5/L6; production email exceptions are not specified yet.
        if level == "L4" or (
            level in ("L5", "L6")
            and computer_rule == "yes"
        ):
            recommended.append("email")

    # Do not preselect desktop alongside laptop.
    # Manager can still select Desktop or any other available asset.
    return recommended

# @login_required
# @require_POST
# def create_joining_kit_request(request, joinee_id):
#     joinee = get_object_or_404(Joinee, pk=joinee_id)

#     if not joinee.employee_code.strip():
#         return HttpResponse(
#             "Enter the Employee Code before creating a Joining Kit request.",
#             status=400,
#         )

#     # Avoid creating another open request for the same joinee.
#     existing = JoiningKitRequest.objects.filter(
#         joinee=joinee,
#         status__in=["pending", "submitted", "in_progress"],
#     ).first()

#     if existing:
#         return redirect(
#             "onboarding:joining_kit_detail",
#             request_id=existing.pk,
#         )

#     kit_request = JoiningKitRequest.objects.create(
#         joinee=joinee,
#         manager_email="abhishek.chuiya@vishakha.com",
#         requested_by=request.user,
#     )

#     return redirect(
#         "onboarding:joining_kit_detail",
#         request_id=kit_request.pk,
#     )


@login_required
@require_POST
def create_joining_kit_request(request, joinee_id):
    require_hr(request.user)

    # The user must explicitly select the Joining Kit trigger.
    if request.POST.get("trigger_joining_kit") != "yes":
        raise PermissionDenied(
            "The Joining Kit trigger was not selected."
        )

    joinee = get_object_or_404(Joinee, pk=joinee_id)

    if not joinee.employee_code.strip():
        return HttpResponse(
            "Generate and save the employee code before creating "
            "a Joining Kit request.",
            status=400,
        )

    existing_request = JoiningKitRequest.objects.filter(
        joinee=joinee,
        status__in=["pending", "submitted", "in_progress"],
    ).first()

    if existing_request:
        return redirect(
            "onboarding:joining_kit_detail",
            request_id=existing_request.pk,
        )

    joining_kit_request = JoiningKitRequest.objects.create(
        joinee=joinee,
        manager_email=MANAGER_EMAIL,
        status="pending",
        requested_by=request.user,
    )

    return redirect(
        "onboarding:joining_kit_detail",
        request_id=joining_kit_request.pk,
    )
@login_required
def manager_joining_kit_dashboard(request):
    if not is_joining_kit_manager(request.user):
        raise PermissionDenied(
            "Only the assigned Joining Kit manager can access this dashboard."
        )

    joining_kit_requests = (
        JoiningKitRequest.objects
        .filter(manager_email__iexact=request.user.email)
        .select_related("joinee", "requested_by")
        .order_by("-created_at")
    )

    return render(
        request,
        "onboarding/manager_joining_kit_dashboard.html",
        {
            "joining_kit_requests": joining_kit_requests,
        },
    )


# @login_required
# def joining_kit_detail(request, request_id):
#     from django.contrib import messages

#     joining_kit_request = get_object_or_404(
#         JoiningKitRequest.objects.select_related(
#             "joinee",
#             "requested_by",
#         ),
#         pk=request_id,
#     )

#     is_manager = is_joining_kit_manager(request.user)
#     is_hr = is_hr_user(request.user)

#     if not (is_manager or is_hr):
#         raise PermissionDenied(
#             "You are not authorized to access this request."
#         )

#     if is_manager:
#         require_assigned_manager(request.user, joining_kit_request)

#     # Only the assigned manager can submit the asset selections.
#     if request.method == "POST":
#         if not is_manager:
#             raise PermissionDenied(
#                 "HR can view the request but cannot submit manager selections."
#             )

#         if joining_kit_request.status != "pending":
#             messages.info(
#                 request,
#                 "This request has already been submitted or processed.",
#             )
#             return redirect(
#                 "onboarding:joining_kit_detail",
#                 request_id=joining_kit_request.pk,
#             )

#         selected_assets = request.POST.getlist("selected_assets")

#         valid_asset_values = {
#             value for value, label in ASSET_OPTIONS
#         }

#         selected_assets = list(dict.fromkeys(
#             asset.strip()
#             for asset in selected_assets
#             if asset.strip() in valid_asset_values
#         ))

#         if not selected_assets:
#             return render(
#                 request,
#                 "onboarding/joining_kit_detail.html",
#                 {
#                     "joining_kit_request": joining_kit_request,
#                     "asset_options": ASSET_OPTIONS,
#                     "recommended_assets": get_recommended_assets(
#                         joining_kit_request.joinee
#                     ),
#                     "can_edit": True,
#                     "is_hr_view": False,
#                     "error": "Please select at least one asset before submitting.",
#                 },
#                 status=400,
#             )

#         joining_kit_request.selected_assets = selected_assets
#         joining_kit_request.required_by = (
#             request.POST.get(
#                 "required_by",
#                 "Before joining date",
#             ).strip()[:100]
#             or "Before joining date"
#         )
#         joining_kit_request.notes = (
#             request.POST.get("notes", "").strip()[:2000]
#         )
#         joining_kit_request.status = "submitted"
#         joining_kit_request.manager_submitted_at = timezone.now()

#         joining_kit_request.save(
#             update_fields=[
#                 "selected_assets",
#                 "required_by",
#                 "notes",
#                 "status",
#                 "manager_submitted_at",
#                 "updated_at",
#             ]
#         )

#         # Create Service Excellence tickets only after a real submission.
#         try:
#             ticket_results = create_joining_kit_tickets(
#                 joining_kit_request
#             )

#             created_count = sum(
#                 ticket.status == "created"
#                 for ticket in ticket_results
#             )
#             failed_count = sum(
#                 ticket.status == "failed"
#                 for ticket in ticket_results
#             )

#             if failed_count:
#                 messages.warning(
#                     request,
#                     f"Joining Kit submitted. {created_count} ticket(s) "
#                     f"created; {failed_count} need attention.",
#                 )
#             else:
#                 messages.success(
#                     request,
#                     f"Joining Kit submitted successfully. "
#                     f"{created_count} ticket(s) created.",
#                 )

#         except Exception:
#             # Keep the submission saved if ticket creation fails.
#             # Do not silently change the request back to pending.
#             messages.warning(
#                 request,
#                 "The asset request was submitted, but ticket creation "
#                 "encountered an error. Please check the server logs.",
#             )

#         return redirect(
#             "onboarding:joining_kit_detail",
#             request_id=joining_kit_request.pk,
#         )

#     # GET requests must only display the page. Never save or redirect here.
#     return render(
#         request,
#         "onboarding/joining_kit_detail.html",
#         {
#             "joining_kit_request": joining_kit_request,
#             "asset_options": ASSET_OPTIONS,
#             "recommended_assets": get_recommended_assets(
#                 joining_kit_request.joinee
#             ),
#             "can_edit": (
#                 is_manager
#                 and joining_kit_request.status == "pending"
#             ),
#             "is_hr_view": is_hr,
#         },
#     )


@login_required
def joining_kit_detail(request, request_id):
    from django.contrib import messages

    joining_kit_request = get_object_or_404(
        JoiningKitRequest.objects.select_related(
            "joinee",
            "requested_by",
        ).prefetch_related("generated_tickets"),
        pk=request_id,
    )

    is_manager = is_joining_kit_manager(request.user)
    is_hr = is_hr_user(request.user)

    if not (is_manager or is_hr):
        raise PermissionDenied(
            "You are not authorized to access this request."
        )

    if is_manager:
        require_assigned_manager(request.user, joining_kit_request)

    if request.method == "POST":
        if not is_manager:
            raise PermissionDenied(
                "Only the assigned manager can submit asset selections."
            )

        if joining_kit_request.status != "pending":
            messages.info(
                request,
                "The asset request has already been submitted.",
            )
            return redirect(
                "onboarding:joining_kit_detail",
                request_id=joining_kit_request.pk,
            )

        selected_assets = request.POST.getlist("selected_assets")
        valid_asset_values = {
            value for value, label in ASSET_OPTIONS
        }

        selected_assets = list(dict.fromkeys(
            asset.strip()
            for asset in selected_assets
            if asset.strip() in valid_asset_values
        ))

        if not selected_assets:
            return render(
                request,
                "onboarding/joining_kit_detail.html",
                {
                    "joining_kit_request": joining_kit_request,
                    "asset_options": ASSET_OPTIONS,
                    "recommended_assets": get_recommended_assets(
                        joining_kit_request.joinee
                    ),
                    "can_edit": True,
                    "is_hr_view": False,
                    "error": (
                        "Please select at least one asset "
                        "before submitting."
                    ),
                    "generated_tickets": (
                        joining_kit_request.generated_tickets.all()
                    ),
                    "can_create_tickets": False,
                },
                status=400,
            )

        joining_kit_request.selected_assets = selected_assets
        joining_kit_request.required_by = (
            request.POST.get(
                "required_by",
                "Before joining date",
            ).strip()[:100]
            or "Before joining date"
        )
        joining_kit_request.notes = (
            request.POST.get("notes", "").strip()[:2000]
        )
        joining_kit_request.status = "submitted"
        joining_kit_request.manager_submitted_at = timezone.now()

        joining_kit_request.save(
            update_fields=[
                "selected_assets",
                "required_by",
                "notes",
                "status",
                "manager_submitted_at",
                "updated_at",
            ]
        )

        messages.success(
            request,
            "Assets submitted successfully. Review the request, "
            "then select Create Tickets when you are ready.",
        )

        return redirect(
            "onboarding:joining_kit_detail",
            request_id=joining_kit_request.pk,
        )

    generated_tickets = list(
        joining_kit_request.generated_tickets.all().order_by("id")
    )

    # Enable ticket creation only after manager submission.
    # Do not expose the action to HR's read-only view.
    can_create_tickets = (
        is_manager
        and joining_kit_request.status in ("submitted", "in_progress")
        and bool(joining_kit_request.selected_assets)
        and any(
            ticket.status in ("pending", "failed")
            for ticket in generated_tickets
        )
        or (
            is_manager
            and joining_kit_request.status in ("submitted", "in_progress")
            and bool(joining_kit_request.selected_assets)
            and not generated_tickets
        )
    )

    return render(
        request,
        "onboarding/joining_kit_detail.html",
        {
            "joining_kit_request": joining_kit_request,
            "asset_options": ASSET_OPTIONS,
            "recommended_assets": get_recommended_assets(
                joining_kit_request.joinee
            ),
            "can_edit": (
                is_manager
                and joining_kit_request.status == "pending"
            ),
            "is_hr_view": is_hr,
            "generated_tickets": generated_tickets,
            "can_create_tickets": can_create_tickets,
        },
    )


@login_required
@require_POST
def create_joining_kit_tickets_view(request, request_id):
    from django.contrib import messages

    joining_kit_request = get_object_or_404(
        JoiningKitRequest.objects.select_related("joinee"),
        pk=request_id,
    )

    require_assigned_manager(request.user, joining_kit_request)

    if joining_kit_request.status not in ("submitted", "in_progress"):
        messages.error(
            request,
            "Submit the asset selections before creating tickets.",
        )
        return redirect(
            "onboarding:joining_kit_detail",
            request_id=joining_kit_request.pk,
        )

    if not joining_kit_request.selected_assets:
        messages.error(
            request,
            "No selected assets were found for this request.",
        )
        return redirect(
            "onboarding:joining_kit_detail",
            request_id=joining_kit_request.pk,
        )

    try:
        ticket_results = create_joining_kit_tickets(
            joining_kit_request
        )

        created_count = sum(
            ticket.status == "created"
            for ticket in ticket_results
        )
        failed_count = sum(
            ticket.status == "failed"
            for ticket in ticket_results
        )
        pending_count = sum(
            ticket.status in ("pending", "creating")
            for ticket in ticket_results
        )

        if failed_count or pending_count:
            messages.warning(
                request,
                f"Ticket processing finished: {created_count} created, "
                f"{failed_count} failed, {pending_count} pending or "
                "requiring investigation. Review the ticket statuses "
                "before retrying.",
            )
        else:
            messages.success(
                request,
                f"Ticket processing finished. {created_count} ticket(s) "
                "are recorded as created.",
            )

        if created_count:
            joining_kit_request.status = "in_progress"
            joining_kit_request.save(
                update_fields=["status", "updated_at"]
            )

    except Exception:
        messages.error(
            request,
            "Ticket creation could not be completed. Check the server "
            "logs and ticket statuses before trying again.",
        )

    return redirect(
        "onboarding:joining_kit_detail",
        request_id=joining_kit_request.pk,
    )
